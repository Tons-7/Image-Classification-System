from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from flask_mail import Mail, Message
import os
from datetime import datetime
import sys
import logging
from PIL import Image
import io
import json
import torch
import tempfile
import shutil
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.append(os.path.join(os.path.dirname(__file__), '..', '..'))

from src.model.predict import EnhancedImagePredictor

app = Flask(__name__)
CORS(app)

app.config['MAIL_SERVER'] = 'smtp.gmail.com'
app.config['MAIL_PORT'] = 587
app.config['MAIL_USE_TLS'] = True
app.config['MAIL_USERNAME'] = os.getenv('EMAIL_USER', '')
app.config['MAIL_PASSWORD'] = os.getenv('EMAIL_PASSWORD', '')
app.config['MAIL_DEFAULT_SENDER'] = os.getenv('EMAIL_USER', '')

mail = Mail(app)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("backend.log"),
        logging.StreamHandler()
    ]
)

MODEL_PATH = os.path.join(os.path.dirname(__file__), '..', 'model', 'output', 'best_model.pth')
CLASS_INFO_PATH = os.path.join(os.path.dirname(__file__), '..', 'model', 'output', 'class_info.json')

try:
    predictor = EnhancedImagePredictor(MODEL_PATH, CLASS_INFO_PATH)
    logging.info("Model loaded successfully")
except Exception as e:
    logging.error(f"Error loading model: {e}")
    predictor = None

CONFIG = {
    'contact_email': 'tonyboutros123987@gmail.com',
    'model_info': {
        'name': 'Image Classification Model',
        'version': '1.0',
        'last_updated': datetime.now().strftime('%Y-%m-%d'),
        'accuracy': '95%',
        'supported_formats': ['jpg', 'jpeg', 'png', 'bmp']
    }
}

CLASSES = [
    {
        'name': 'bear',
        'display_name': 'Bear',
        'description': 'Different species of bears',
        'examples': ['Brown bear', 'Black bear', 'Polar bear']
    },
    {
        'name': 'butterfly',
        'display_name': 'Butterfly',
        'description': 'Various species of butterflies',
        'examples': ['Monarch', 'Swallowtail', 'Blue morpho']
    },
    {
        'name': 'camel',
        'display_name': 'Camel',
        'description': 'Different types of camels',
        'examples': ['Dromedary', 'Bactrian', 'Wild camel']
    },
    {
        'name': 'capybara',
        'display_name': 'Capybara',
        'description': 'The world\'s largest rodent',
        'examples': ['Capybara', 'Giant capybara']
    },
    {
        'name': 'cat',
        'display_name': 'Cat',
        'description': 'Various breeds of domestic cats',
        'examples': ['Persian', 'Siamese', 'Maine Coon']
    },
    {
        'name': 'chicken',
        'display_name': 'Chicken',
        'description': 'Different breeds of chickens',
        'examples': ['Leghorn', 'Rhode Island Red', 'Plymouth Rock']
    },
    {
        'name': 'cow',
        'display_name': 'Cow',
        'description': 'Various breeds of cattle',
        'examples': ['Holstein', 'Jersey', 'Angus']
    },
    {
        'name': 'dog',
        'display_name': 'Dog',
        'description': 'Different breeds of dogs',
        'examples': ['Labrador', 'German Shepherd', 'Golden Retriever']
    },
    {
        'name': 'elephant',
        'display_name': 'Elephant',
        'description': 'Different species of elephants',
        'examples': ['African elephant', 'Asian elephant']
    },
    {
        'name': 'horse',
        'display_name': 'Horse',
        'description': 'Various breeds of horses',
        'examples': ['Arabian', 'Thoroughbred', 'Mustang']
    },
    {
        'name': 'sheep',
        'display_name': 'Sheep',
        'description': 'Different breeds of sheep',
        'examples': ['Merino', 'Suffolk', 'Hampshire']
    },
    {
        'name': 'spider',
        'display_name': 'Spider',
        'description': 'Various species of spiders',
        'examples': ['Tarantula', 'Black widow', 'Wolf spider']
    },
    {
        'name': 'squirrel',
        'display_name': 'Squirrel',
        'description': 'Different types of squirrels',
        'examples': ['Gray squirrel', 'Red squirrel', 'Flying squirrel']
    }
]


def add_heading_with_border(document, text, level=1):
    heading = document.add_heading(text, level=level)
    p = heading._element
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    pPr.append(pBdr)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '4472C4')
    pBdr.append(bottom)
    return heading


def create_results_document(results):
    doc = Document()

    doc.core_properties.title = "Image Classification Results"
    doc.core_properties.author = "AI Image Classifier"

    title = doc.add_heading('Image Classification Results', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    doc.add_paragraph(f'Generated on: {timestamp}')
    doc.add_paragraph(f'Total Images Processed: {len(results)}')
    doc.add_paragraph()

    for result in results:
        add_heading_with_border(doc, f"Results for {result['filename']}", level=1)

        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'

        header_cells = table.rows[0].cells
        header_cells[0].text = 'Rank'
        header_cells[1].text = 'Class'
        header_cells[2].text = 'Confidence'

        for pred in result['predictions']:
            row_cells = table.add_row().cells
            row_cells[0].text = str(pred['rank'])
            row_cells[1].text = pred['class_name']
            confidence = f"{pred['probability'] * 100:.2f}%"

            if pred['probability'] > 0.9:
                confidence += " (Very High)"
            elif pred['probability'] > 0.7:
                confidence += " (High)"
            elif pred['probability'] > 0.5:
                confidence += " (Moderate)"
            else:
                confidence += " (Low)"

            row_cells[2].text = confidence

        top_pred = result['predictions'][0]
        doc.add_paragraph()
        summary = doc.add_paragraph()
        summary.add_run('Top Prediction: ').bold = True
        summary.add_run(f"{top_pred['class_name']} with {top_pred['probability'] * 100:.2f}% confidence")

        doc.add_paragraph()
        doc.add_paragraph()

    return doc


@app.route('/api/classes', methods=['GET'])
def get_classes():
    try:
        if os.path.exists(CLASS_INFO_PATH):
            with open(CLASS_INFO_PATH, 'r') as f:
                class_info = json.load(f)
            return jsonify({
                'success': True,
                'data': {
                    'classes': class_info,
                    'model_info': CONFIG['model_info']
                }
            })
        else:
            return jsonify({
                'success': True,
                'data': {
                    'classes': CLASSES,
                    'model_info': CONFIG['model_info']
                }
            })
    except Exception as e:
        logging.error(f"Error getting classes: {e}")
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@app.route('/api/config', methods=['GET'])
def get_config():
    return jsonify({
        'success': True,
        'data': {
            'contact_email': CONFIG['contact_email']
        }
    })


@app.route('/api/contact', methods=['POST'])
def contact():
    try:
        data = request.json
        if not data:
            return jsonify({
                'success': False,
                'error': 'No data provided'
            }), 400

        is_class_request = 'class_name' in data and 'description' in data and 'examples' in data

        if is_class_request:
            required_fields = ['name', 'email', 'class_name', 'description', 'examples']
        else:
            required_fields = ['name', 'email', 'message']

        for field in required_fields:
            if field not in data:
                return jsonify({
                    'success': False,
                    'error': f'Missing required field: {field}'
                }), 400

        sender_email = data['email']

        if is_class_request:
            subject = f'New Class Request: {data["class_name"]}'
            body = f'''
            Name: {data['name']}
            Email: {data['email']}
            Class Name: {data['class_name']}
            Description: {data['description']}
            Examples: {data['examples']}
            '''
            html = f'''
            <!DOCTYPE html>
            <html>
            <head>
                <style>
                    body {{ font-family: Arial, sans-serif; line-height: 1.6; color: #333; }}
                    .container {{ max-width: 600px; margin: 0 auto; padding: 20px; }}
                    .header {{ background: linear-gradient(135deg, #1a237e, #0d47a1); color: white; padding: 30px; text-align: center; border-radius: 10px 10px 0 0; }}
                    .content {{ background: #ffffff; padding: 30px; border: 1px solid #e0e0e0; border-radius: 0 0 10px 10px; }}
                    .section {{ margin-bottom: 20px; padding: 15px; background: #f8f9fa; border-radius: 5px; }}
                    .label {{ font-weight: bold; color: #1a237e; }}
                    .footer {{ text-align: center; margin-top: 20px; padding-top: 20px; border-top: 1px solid #e0e0e0; color: #666; }}
                </style>
            </head>
            <body>
                <div class="container">
                    <div class="header">
                        <h1>New Class Request</h1>
                        <p>AI Image Classification System</p>
                    </div>
                    <div class="content">
                        <div class="section">
                            <p><span class="label">Name:</span> {data['name']}</p>
                            <p><span class="label">Email:</span> {data['email']}</p>
                        </div>
                        <div class="section">
                            <p><span class="label">Class Name:</span> {data['class_name']}</p>
                            <p><span class="label">Description:</span> {data['description']}</p>
                            <p><span class="label">Examples:</span> {data['examples']}</p>
                        </div>
                        <div class="footer">
                            <p>This is an automated message from the AI Image Classification System</p>
                            <p>Timestamp: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
                        </div>
                    </div>
                </div>
            </body>
            </html>
            '''
        else:
            subject = f'New Contact Form Submission from {data["name"]}'
            body = f'''
            Name: {data['name']}
            Email: {data['email']}
            Message: {data['message']}
            '''
            html = f'''
            <!DOCTYPE html>
            <html>
            <head>
                <style>
                    body {{ font-family: Arial, sans-serif; line-height: 1.6; color: #333; }}
                    .container {{ max-width: 600px; margin: 0 auto; padding: 20px; }}
                    .header {{ background: linear-gradient(135deg, #1a237e, #0d47a1); color: white; padding: 30px; text-align: center; border-radius: 10px 10px 0 0; }}
                    .content {{ background: #ffffff; padding: 30px; border: 1px solid #e0e0e0; border-radius: 0 0 10px 10px; }}
                    .section {{ margin-bottom: 20px; padding: 15px; background: #f8f9fa; border-radius: 5px; }}
                    .label {{ font-weight: bold; color: #1a237e; }}
                    .footer {{ text-align: center; margin-top: 20px; padding-top: 20px; border-top: 1px solid #e0e0e0; color: #666; }}
                </style>
            </head>
            <body>
                <div class="container">
                    <div class="header">
                        <h1>New Contact Form Submission</h1>
                        <p>AI Image Classification System</p>
                    </div>
                    <div class="content">
                        <div class="section">
                            <p><span class="label">Name:</span> {data['name']}</p>
                            <p><span class="label">Email:</span> {data['email']}</p>
                        </div>
                        <div class="section">
                            <p><span class="label">Message:</span></p>
                            <p>{data['message']}</p>
                        </div>
                        <div class="footer">
                            <p>This is an automated message from the AI Image Classification System</p>
                            <p>Timestamp: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
                        </div>
                    </div>
                </div>
            </body>
            </html>
            '''

        msg = Message(
            subject=subject,
            recipients=[CONFIG['contact_email']],
            sender=sender_email,
            body=body,
            html=html
        )

        mail.send(msg)

        confirmation_subject = 'Thank you for your request' if is_class_request else 'Thank you for contacting us'
        confirmation_body = f'''
        Dear {data['name']},

        Thank you for your request. We have received it and will review it soon.

        Best regards,
        The Team
        '''
        confirmation_html = f'''
        <!DOCTYPE html>
        <html>
        <head>
            <style>
                body {{ font-family: Arial, sans-serif; line-height: 1.6; color: #333; }}
                .container {{ max-width: 600px; margin: 0 auto; padding: 20px; }}
                .header {{ background: linear-gradient(135deg, #1a237e, #0d47a1); color: white; padding: 30px; text-align: center; border-radius: 10px 10px 0 0; }}
                .content {{ background: #ffffff; padding: 30px; border: 1px solid #e0e0e0; border-radius: 0 0 10px 10px; }}
                .message {{ margin: 20px 0; padding: 20px; background: #f8f9fa; border-radius: 5px; }}
                .footer {{ text-align: center; margin-top: 20px; padding-top: 20px; border-top: 1px solid #e0e0e0; color: #666; }}
                .button {{ display: inline-block; padding: 12px 24px; background: linear-gradient(135deg, #1a237e, #0d47a1); color: white; text-decoration: none; border-radius: 5px; margin-top: 20px; }}
            </style>
        </head>
        <body>
            <div class="container">
                <div class="header">
                    <h1>Thank You for Your Request</h1>
                    <p>AI Image Classification System</p>
                </div>
                <div class="content">
                    <div class="message">
                        <p>Dear {data['name']},</p>
                        <p>Thank you for your interest in our AI Image Classification System. We have received your request and will review it shortly.</p>
                        <p>Our team will carefully evaluate your submission and get back to you as soon as possible.</p>
                    </div>
                    <div style="text-align: center;">
                        <a href="http://localhost:3000" class="button">Return to Website</a>
                    </div>
                    <div class="footer">
                        <p>This is an automated message from the AI Image Classification System</p>
                        <p>If you have any questions, please don't hesitate to contact us.</p>
                        <p>Timestamp: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</p>
                    </div>
                </div>
            </div>
        </body>
        </html>
        '''

        confirmation_msg = Message(
            subject=confirmation_subject,
            recipients=[data['email']],
            sender=CONFIG['contact_email'],
            body=confirmation_body,
            html=confirmation_html
        )
        mail.send(confirmation_msg)

        return jsonify({
            'success': True,
            'message': 'Request submitted successfully'
        })

    except Exception as e:
        logging.error(f"Error sending email: {e}")
        return jsonify({
            'success': False,
            'error': 'Failed to send email. Please try again later.'
        }), 500


@app.route('/api/predict', methods=['POST'])
def predict():
    if 'file' not in request.files:
        return jsonify({
            'success': False,
            'error': 'No file provided'
        }), 400

    file = request.files['file']
    if file.filename == '':
        return jsonify({
            'success': False,
            'error': 'No file selected'
        }), 400

    try:
        temp_path = os.path.join(os.path.dirname(__file__), 'temp_image.jpg')
        file.save(temp_path)

        predictions = predictor.predict(temp_path, top_k=5)

        os.remove(temp_path)

        if not predictions:
            return jsonify({
                'success': False,
                'error': 'Failed to process image. Please try a different image.'
            }), 400

        return jsonify({
            'success': True,
            'predictions': predictions
        })

    except Exception as e:
        logging.error(f"Error processing image: {e}")
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@app.route('/api/predict-batch', methods=['POST'])
def predict_batch():
    if 'files[]' not in request.files:
        return jsonify({'success': False, 'error': 'No files provided'}), 400

    files = request.files.getlist('files[]')
    if not files:
        return jsonify({'success': False, 'error': 'No files selected'}), 400

    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            results = []

            for file in files:
                if file.filename:
                    temp_path = os.path.join(temp_dir, file.filename)
                    file.save(temp_path)

                    predictions = predictor.predict(temp_path)

                    results.append({
                        'filename': file.filename,
                        'predictions': predictions
                    })

            doc = create_results_document(results)

            temp_docx = os.path.join(temp_dir, 'classification_results.docx')
            doc.save(temp_docx)

            with open(temp_docx, 'rb') as f:
                docx_data = f.read()

            return send_file(
                io.BytesIO(docx_data),
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                as_attachment=True,
                download_name=f'classification_results_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
            )

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


if __name__ == '__main__':
    app.run(port=5001, debug=True)
